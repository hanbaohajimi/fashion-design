from __future__ import annotations

import json
import os
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
DOC_PATTERN = os.environ.get("FLOWBOOK_DOC_PATTERN", "演示视频*_拍摄流程书.docx")
DOCS = sorted((ROOT / "doc").glob(DOC_PATTERN))
REQUIRED = {
    "01": ["冬季长款双面羊毛大衣", "枪驳领", "局部重绘", "版本时间线", "系列 DNA", "一致性", "面料库严格只有图片和名称", "图片为唯一材质参考"],
    "02": ["夏季连衣裙", "方领", "完成并引用", "安全橡皮", "廓形参考", "版本时间线", "V2", "面料库严格只有图片和名称", "图片为材质参考"],
    "03": ["雪纺衫", "飘带领", "趋势分析", "避用", "最多 16 张", "系列计划", "相似度矩阵", "面料库回流结束时只有图片和名称", "图片为材质参考"],
}

FORBIDDEN_FABRIC_LIBRARY_CLAIMS = (
    "品类/成分/GSM/季节",
    "面料档案登记",
    "可验证物性跟着面料卡",
    "面料卡带来 92/8",
    "确认档案中的棉 68%",
    "门幅/弹力/透明度编辑",
    "FABRIC SPEC",
    "PROJECT MATERIAL NOTE",
    "普通白板",
    "项目材料约束",
    "GSM",
    "门幅",
    "羊毛90%",
    "羊毛 90%",
    "羊绒10%",
    "羊绒 10%",
    "棉72%",
    "棉 72%",
    "真丝28%",
    "真丝 28%",
    "真丝100%",
    "真丝 100%",
    "项目材料",
    "普通便签",
    "普通项目",
    "项目备注",
    "材料备注",
)

FORBIDDEN_OLD_GARMENTS = {
    "01": ("短款雕塑感通勤夹克", "深海斜纹通勤系列"),
    "02": ("模块风衣裙", "风衣裙", "直筒长袖"),
    "03": ("Urban Shelter", "宽肩短外套", "城市防护外套", "战术背心"),
}


def table_width_audit(doc: Document) -> list[str]:
    problems: list[str] = []
    for index, table in enumerate(doc.tables, 1):
        tbl_pr = table._tbl.tblPr
        width = tbl_pr.find(qn("w:tblW"))
        indent = tbl_pr.find(qn("w:tblInd"))
        grid = [int(col.get(qn("w:w"))) for col in table._tbl.tblGrid]
        if width is None or width.get(qn("w:type")) != "dxa" or int(width.get(qn("w:w"))) != 9360:
            problems.append(f"table {index}: tblW != 9360")
        if indent is None or indent.get(qn("w:type")) != "dxa" or int(indent.get(qn("w:w"))) != 120:
            problems.append(f"table {index}: tblInd != 120")
        if not grid or sum(grid) != 9360:
            problems.append(f"table {index}: grid sum {sum(grid)}")
    return problems


def validate(path: Path) -> dict[str, object]:
    doc = Document(path)
    table_text = [cell.text for table in doc.tables for row in table.rows for cell in row.cells]
    full_text = "\n".join([*(p.text for p in doc.paragraphs), *table_text])
    number = re.search(r"演示视频(\d{2})_", path.name).group(1)
    errors: list[str] = []
    for phrase in REQUIRED[number]:
        if phrase not in full_text:
            errors.append(f"missing phrase: {phrase}")
    for phrase in ("面料库记录", "库内不填写、不显示、不携带任何其他面料信息", "关闭面料库"):
        if phrase not in full_text:
            errors.append(f"missing strict fabric-library boundary: {phrase}")
    for phrase in FORBIDDEN_FABRIC_LIBRARY_CLAIMS:
        if phrase in full_text:
            errors.append(f"outdated fabric-library claim: {phrase}")
    for phrase in FORBIDDEN_OLD_GARMENTS[number]:
        if phrase in full_text:
            errors.append(f"outdated garment theme: {phrase}")
    for placeholder in ("TODO", "TBD", "[[", ":codex-", "待补充"):
        if placeholder in full_text:
            errors.append(f"placeholder token: {placeholder}")
    fake_lists = [
        p.text
        for p in doc.paragraphs
        if not p.style.name.startswith("Heading")
        and re.match(r"^\s*(?:•|-|\d+[.)])\s+", p.text)
    ]
    if fake_lists:
        errors.append(f"fake list paragraphs: {len(fake_lists)}")
    scene_headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 2" and re.match(r"^\d{2}\s", p.text)]
    if len(scene_headings) < 10:
        errors.append(f"only {len(scene_headings)} scene headings")
    numbered_paras = sum(1 for p in doc.paragraphs if p._p.pPr is not None and p._p.pPr.numPr is not None)
    if numbered_paras < 30:
        errors.append(f"too few real list paragraphs: {numbered_paras}")
    if len(doc.tables) < 10:
        errors.append(f"too few structured tables: {len(doc.tables)}")
    section = doc.sections[0]
    if round(section.page_width.inches, 2) != 8.5 or round(section.page_height.inches, 2) != 11.0:
        errors.append("page size is not US Letter")
    margins = [section.top_margin.inches, section.right_margin.inches, section.bottom_margin.inches, section.left_margin.inches]
    if any(abs(value - 1.0) > 0.01 for value in margins):
        errors.append(f"incorrect margins: {margins}")
    errors.extend(table_width_audit(doc))
    return {
        "file": path.name,
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "scene_headings": len(scene_headings),
        "numbered_paragraphs": numbered_paras,
        "characters": len(full_text),
        "errors": errors,
    }


if len(DOCS) != 3:
    raise SystemExit(f"expected 3 flowbooks, found {len(DOCS)}")
reports = [validate(path) for path in DOCS]
print(json.dumps(reports, ensure_ascii=False, indent=2))
if any(report["errors"] for report in reports):
    raise SystemExit(1)
